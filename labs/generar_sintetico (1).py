"""
Generador de base de datos sintética de consumo eléctrico por subestación (kWh por hora).

Resultado: una tabla con las columnas
    Medidor | Fecha | Hora del día | Energía activa (kWh)
con 26 medidores (subestaciones) x 5 años de datos horarios = 1.138.800 filas.
Todas las fechas son pasadas: van de 2021-09-01 a 2026-08-30 (último día del Excel real).

Idea general:
  1. Se lee la hoja 'demanda' del Excel original (matriz: fechas en columnas,
     horas 0-23 en filas) y se pasa a formato "largo": una fila por fecha+hora.
  2. Para cada FECHA se escoge un día real "plantilla" (la forma de la curva de 24 h):
       - si la fecha existe en el Excel real, se usa ese mismo día real;
       - si no, se elige al azar un día real con el mismo mes y día de la semana.
     La plantilla se comparte entre todas las subestaciones de esa fecha, porque
     pertenecen a la misma institución: si un día es festivo o de vacaciones,
     todas las subestaciones consumen menos ese día.
  3. Cada subestación (medidor) tiene su propia "personalidad", fija y aleatoria:
       - escala: qué tan grande es (consume más o menos que el medidor original),
       - intensidad del pico: cuánto sube en horario laboral respecto a la carga base,
       - tendencia anual: crece o decrece un poco año a año.
  4. Encima se aplica ruido independiente por medidor (por día y por hora), de modo
     que dos subestaciones nunca tienen exactamente la misma curva.
  5. Se redondea a múltiplos de 0,25 kWh, igual que la resolución del medidor real.

Uso:
  python generar_sintetico.py <excel_origen> <carpeta_salida>
"""
import sys
import numpy as np
import pandas as pd
import openpyxl

ORIGEN, SALIDA = sys.argv[1], sys.argv[2]

COL_MED, COL_FECHA, COL_HORA, COL_KWH = "Medidor", "Fecha", "Hora del día", "Energía activa (kWh)"
SEMILLA = 42                     # semilla fija: si se vuelve a correr, sale exactamente lo mismo
N_MEDIDORES = 26
FECHA_INICIO, FECHA_FIN = "2021-09-01", "2026-08-30"
SIGMA_DIA = 0.04                 # ±4 % de variación típica del total de un día
SIGMA_HORA = 0.025               # ±2,5 % de variación típica entre horas

rng = np.random.default_rng(SEMILLA)

# ---------- 1. Leer la matriz original y pasarla a formato largo ----------
ws = openpyxl.load_workbook(ORIGEN, data_only=True)["demanda"]
registros = []
for col in range(1, ws.max_column + 1):
    fecha = ws.cell(2, col).value            # fila 2 = fechas
    if fecha is None:
        continue                             # columnas vacías (por celdas combinadas)
    for hora in range(24):                   # filas 3..26 = horas 0..23
        registros.append((fecha, hora, ws.cell(3 + hora, col).value))

real = pd.DataFrame(registros, columns=[COL_FECHA, COL_HORA, COL_KWH])
real[COL_FECHA] = pd.to_datetime(real[COL_FECHA].astype(str), dayfirst=True)
real[COL_KWH] = pd.to_numeric(real[COL_KWH], errors="coerce")

# Matriz días x 24 horas; los días con horas vacías no sirven de plantilla
matriz = real.pivot(index=COL_FECHA, columns=COL_HORA, values=COL_KWH).dropna()
plantillas = matriz.to_numpy()                                    # (n_dias_reales, 24)
clave_real = (matriz.index.month * 10 + matriz.index.dayofweek).to_numpy()
grupos = pd.Series(range(len(matriz))).groupby(clave_real).apply(list).to_dict()
posicion_real = {f: i for i, f in enumerate(matriz.index)}       # fecha real -> fila

# ---------- 2. Calendario y plantilla de cada fecha ----------
fechas = pd.date_range(FECHA_INICIO, FECHA_FIN, freq="D")
n_dias = len(fechas)
claves = (fechas.month * 10 + fechas.dayofweek).to_numpy()
idx = np.array([
    posicion_real[f] if f in posicion_real else rng.choice(grupos[c])
    for f, c in zip(fechas, claves)
])
base = plantillas[idx]                                            # (n_dias, 24)
carga_base = np.percentile(plantillas, 10)                        # consumo "de madrugada"
anios = ((fechas - fechas[-1]).days / 365.25).to_numpy()[:, None] # 0 al final, negativo antes

# ---------- 3. Personalidad de cada medidor ----------
medidores = pd.DataFrame({
    COL_MED: [f"SUB-{i:02d}" for i in range(1, N_MEDIDORES + 1)],
    "Escala": np.round(rng.lognormal(mean=np.log(0.6), sigma=0.6, size=N_MEDIDORES), 3),
    "Intensidad_pico": np.round(rng.uniform(0.6, 1.4, size=N_MEDIDORES), 3),
    "Tendencia_anual": np.round(rng.normal(0.02, 0.03, size=N_MEDIDORES), 4),
})

# ---------- 4. Generar los datos medidor por medidor ----------
partes = []
for _, m in medidores.iterrows():
    # La intensidad del pico estira o encoge la parte que está por encima de la carga base
    curva = carga_base + m.Intensidad_pico * (base - carga_base)
    curva = curva * m.Escala * (1 + m.Tendencia_anual) ** anios
    factor_dia = rng.normal(1.0, SIGMA_DIA, size=(n_dias, 1))
    ruido_hora = rng.normal(1.0, SIGMA_HORA, size=(n_dias, 24))
    kwh = np.round(np.clip(curva * factor_dia * ruido_hora, 0, None) * 4) / 4
    partes.append(pd.DataFrame({
        COL_MED: m[COL_MED],
        COL_FECHA: np.repeat(fechas.strftime("%Y-%m-%d"), 24),
        COL_HORA: np.tile(np.arange(24), n_dias),
        COL_KWH: kwh.ravel(),
    }))

sint = pd.concat(partes, ignore_index=True)
sint.to_csv(f"{SALIDA}/consumo_sintetico_subestaciones.csv", index=False, encoding="utf-8-sig")
medidores.to_csv(f"{SALIDA}/medidores.csv", index=False, encoding="utf-8-sig")
print(f"Filas: {len(sint):,} | Medidores: {N_MEDIDORES} | {fechas[0].date()} -> {fechas[-1].date()}")
